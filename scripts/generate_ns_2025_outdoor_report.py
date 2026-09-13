"""Generate the live SGCS Normal Staff 2025 OUTDOOR reconciliation report."""

from __future__ import annotations

from collections import defaultdict
from contextlib import closing
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from zoneinfo import ZoneInfo

from wecare_spectacle import _connect


OUTPUT = Path("docs/reports/SGCS-NS-2025-Outdoor-Claims-Limits.xlsx")
TAG = "[MIGRATION:SGCS-NS-2025-OUTDOOR-R1]"
LIMIT_ID = 7
GLOBAL_LIMIT = Decimal("9000.00")
HISTORICAL_OVERRIDES = {"206"}

EXPECTED_SOURCE = {
    "102": Decimal("8900.00"),
    "140": Decimal("9000.00"),
    "162": Decimal("9000.00"),
    "170": Decimal("6600.00"),
    "176": Decimal("1523.00"),
    "193": Decimal("7800.00"),
    "194": Decimal("5606.00"),
    "199": Decimal("9000.00"),
    "200": Decimal("9000.00"),
    "206": Decimal("2000.00"),
    "208": Decimal("9000.00"),
    "211": Decimal("5940.00"),
    "237": Decimal("9000.00"),
    "238": Decimal("3402.00"),
    "239": Decimal("0.00"),
    "244": Decimal("8017.00"),
    "247": Decimal("9000.00"),
    "248": Decimal("0.00"),
}

AUDIT_ACTIONS = {
    469: ("UPDATED_AND_APPROVED", "UNDER_REVIEW claim corrected to 1,523 and approved; quarter ID 64 assigned."),
    388: ("SPLIT_EXISTING_OTHER", "Old 7,800 aggregate changed to its 2,800 OTHER component; quarter ID 62 assigned."),
    646: ("CREATED_DENTAL_SPLIT", "5,000 DENTAL component created; annual DENTAL quarter ID 69 assigned."),
    647: ("CREATED_MISSING", "Missing 2,400 OTHER claim created using the confirmed 2025-12-13 bill date."),
    648: ("CREATED_HISTORICAL_NS_OVERRIDE", "Historical NS 2025 OUTDOOR claim created for EPF 206."),
    649: ("CREATED_MISSING", "Missing 2,650 OTHER balance-sheet claim created for EPF 237."),
    650: ("CREATED_MISSING", "Missing 994 OTHER balance-sheet claim created for EPF 237."),
    396: ("REVIEWED_CONFIRMED_OTHER", "Existing 8,017 claim reviewed and explicitly confirmed as OTHER."),
    434: ("UPDATED_AND_APPROVED", "Existing 9,000 UNDER_REVIEW claim approved without inserting a duplicate."),
}


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
GOOD_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE_FILL = PatternFill("solid", fgColor="17365D")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="B7B7B7")
MONEY = '#,##0.00'


def query_all(cur, sql: str, params=()):
    cur.execute(sql, params)
    return cur.fetchall()


def active_cohort(cur):
    rows = query_all(
        cur,
        """
        SELECT upd.epf_no,
               TRIM(CONCAT_WS(' ', NULLIF(upd.first_name,''), NULLIF(upd.last_name,''))) employee_name,
               upd.user_status, ucd.staff_category, ucd.insurance_policy,
               ucd.company_type, au.id user_id
        FROM user_personal_details upd
        JOIN user_company_details ucd ON ucd.id=upd.user_company_details
        JOIN application_user au ON au.user_personal_details=upd.id
        WHERE upd.user_status='ACTIVE' AND ucd.company_type='SGCS'
          AND ucd.staff_category='NS' AND ucd.insurance_policy='P-NORMAL'
        ORDER BY CAST(upd.epf_no AS UNSIGNED)
        """,
    )
    by_epf = {str(row["epf_no"]): row for row in rows}
    for epf in HISTORICAL_OVERRIDES:
        if epf not in by_epf:
            override = query_all(
                cur,
                """
                SELECT upd.epf_no,
                       TRIM(CONCAT_WS(' ', NULLIF(upd.first_name,''), NULLIF(upd.last_name,''))) employee_name,
                       upd.user_status, ucd.staff_category, ucd.insurance_policy,
                       ucd.company_type, au.id user_id
                FROM user_personal_details upd
                JOIN user_company_details ucd ON ucd.id=upd.user_company_details
                JOIN application_user au ON au.user_personal_details=upd.id
                WHERE upd.user_status='ACTIVE' AND ucd.company_type='SGCS'
                  AND upd.epf_no=%s
                """,
                (epf,),
            )
            if len(override) != 1:
                raise RuntimeError(f"Cannot resolve historical override EPF {epf}: {override}")
            by_epf[epf] = override[0]
    return sorted(by_epf.values(), key=lambda row: int(row["epf_no"]))


def claim_rows(cur, user_ids):
    marks = ",".join(["%s"] * len(user_ids))
    return query_all(
        cur,
        f"""
        SELECT upd.epf_no, cr.id claim_id, cr.request_id, cr.request_status,
               cr.request_amount, cr.approved_amount,
               icd.treatment_category, icd.from_treatment_date,
               icd.to_treatment_date, cr.created_date,
               cr.insurance_details_limit_id limit_id,
               cr.insurance_quarter_id quarter_id, cr.remark
        FROM claims_request cr
        JOIN application_user au ON au.id=cr.employee
        JOIN user_personal_details upd ON upd.id=au.user_personal_details
        JOIN insurance_claims_details icd ON icd.id=cr.insurance_claims_details
        WHERE cr.employee IN ({marks}) AND cr.insurance_details_limit_id=%s
        ORDER BY CAST(upd.epf_no AS UNSIGNED), cr.id
        """,
        (*user_ids, LIMIT_ID),
    )


def limit_rows(cur):
    global_row = query_all(
        cur,
        """
        SELECT id limit_id, insurance_staff_category_period period_id,
               insurance_policy, treatment, global_limit
        FROM insurance_details_limit WHERE id=%s
        """,
        (LIMIT_ID,),
    )
    if len(global_row) != 1:
        raise RuntimeError("OUTDOOR limit ID 7 is missing")
    quarters = query_all(
        cur,
        """
        SELECT iq.id quarter_id, iq.treatment_category_code,
               iq.quarter_limit, iq.from_date, iq.to_date
        FROM insurance_quarter iq
        WHERE iq.insurance_details_id=%s
        ORDER BY iq.treatment_category_code, iq.from_date, iq.id
        """,
        (LIMIT_ID,),
    )
    return global_row[0], quarters


def style_table(ws, name):
    if ws.max_row < 2:
        return
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ref
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32


def fit_columns(ws, maximum=48):
    for idx in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, idx).value or "") for row in range(1, ws.max_row + 1)]
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(v) for v in values) + 2, maximum)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)


def add_summary(wb, cohort, claims, marker_count, generated_at):
    ws = wb.active
    ws.title = "Summary"
    ws.merge_cells("A1:F1")
    ws["A1"] = "SGCS Normal Staff 2025 Outdoor Claims and Limits"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    approved_total = sum(
        (row["approved_amount"] or Decimal(0))
        for row in claims if row["request_status"] == "APPROVED"
    )
    rows = [
        ("Database snapshot", "sgcs_care (live)"),
        ("Generated at", generated_at),
        ("Policy period", "2025-01-21 to 2026-01-20"),
        ("Staff category / policy", "NS / P-NORMAL"),
        ("Treatment", "OUTDOOR"),
        ("Global outdoor limit per employee", GLOBAL_LIMIT),
        ("Cohort employees", len(cohort)),
        ("Current NS / P-NORMAL employees", len(cohort) - len(HISTORICAL_OVERRIDES)),
        ("Historical NS override employees", len(HISTORICAL_OVERRIDES)),
        ("Approved claim records", sum(r["request_status"] == "APPROVED" for r in claims)),
        ("Under-review claim records", sum(r["request_status"] == "UNDER_REVIEW" for r in claims)),
        ("Rejected claim records", sum(r["request_status"] == "REJECTED" for r in claims)),
        ("Approved amount across cohort", approved_total),
        ("Tagged migration/reconciliation claims", marker_count),
        ("Migration marker", TAG),
        ("Utilization rule", "Only APPROVED claims and approved_amount reduce the available limit."),
        ("Duplicate rule", "Existing live claims were reconciled; no matching claim was inserted twice."),
    ]
    for row_no, (label, value) in enumerate(rows, 3):
        ws.cell(row_no, 1, label).font = Font(bold=True)
        ws.cell(row_no, 1).fill = SUBHEADER_FILL
        ws.merge_cells(start_row=row_no, start_column=2, end_row=row_no, end_column=6)
        ws.cell(row_no, 2, value)
        if isinstance(value, Decimal):
            ws.cell(row_no, 2).number_format = MONEY
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 75
    for col in "CDEF":
        ws.column_dimensions[col].width = 3
    ws.freeze_panes = "A3"


def add_employee_limits(wb, cohort, claims):
    ws = wb.create_sheet("Employee Category Limits")
    ws.append([
        "EPF", "Employee Name", "User Status", "Current Staff", "Current Policy",
        "Cohort Basis", "Approved OTHER", "Approved DENTAL",
        "Total Approved OUTDOOR", "Global OUTDOOR Limit", "Global OUTDOOR Balance",
        "Approved Claims", "Under Review Claims", "Under Review Requested",
        "Rejected Claims", "Total Claim Records", "Migration-tagged Claims", "Notes",
    ])
    by_epf = defaultdict(list)
    for row in claims:
        by_epf[str(row["epf_no"])].append(row)
    for employee in cohort:
        epf = str(employee["epf_no"])
        records = by_epf[epf]
        approved = [r for r in records if r["request_status"] == "APPROVED"]
        other = sum((r["approved_amount"] or Decimal(0)) for r in approved if r["treatment_category"] == "OTHER")
        dental = sum((r["approved_amount"] or Decimal(0)) for r in approved if r["treatment_category"] == "DENTAL")
        total = sum((r["approved_amount"] or Decimal(0)) for r in approved)
        under_review = [r for r in records if r["request_status"] == "UNDER_REVIEW"]
        rejected = [r for r in records if r["request_status"] == "REJECTED"]
        tagged = [r for r in records if TAG in (r["remark"] or "")]
        basis = "HISTORICAL_NS_2025_OVERRIDE" if epf in HISTORICAL_OVERRIDES else "CURRENT_NS_P_NORMAL"
        notes = "Current profile is not NS/P-NORMAL; included because an approved historical NS 2025 claim was confirmed." if epf in HISTORICAL_OVERRIDES else None
        ws.append([
            epf, employee["employee_name"], employee["user_status"],
            employee["staff_category"], employee["insurance_policy"], basis,
            other, dental, total, GLOBAL_LIMIT, GLOBAL_LIMIT - total,
            len(approved), len(under_review),
            sum((r["request_amount"] or Decimal(0)) for r in under_review),
            len(rejected), len(records), len(tagged), notes,
        ])
        row_no = ws.max_row
        for col in range(7, 12):
            ws.cell(row_no, col).number_format = MONEY
        ws.cell(row_no, 14).number_format = MONEY
        if total > GLOBAL_LIMIT:
            ws.cell(row_no, 9).fill = WARN_FILL
            ws.cell(row_no, 11).fill = WARN_FILL
        elif epf in EXPECTED_SOURCE:
            ws.cell(row_no, 9).fill = GOOD_FILL
            ws.cell(row_no, 11).fill = GOOD_FILL
    style_table(ws, "OutdoorEmployeeLimits")
    fit_columns(ws)


def add_claim_details(wb, claims):
    ws = wb.create_sheet("Claim Details")
    ws.append([
        "EPF", "Claim ID", "Request ID", "Status", "Treatment Category",
        "Requested Amount", "Approved Amount", "From Treatment Date",
        "To Treatment Date", "Created Date", "Limit ID", "Quarter ID",
        "Migration Marker", "Remark",
    ])
    for row in claims:
        ws.append([
            str(row["epf_no"]), row["claim_id"], row["request_id"],
            row["request_status"], row["treatment_category"],
            row["request_amount"], row["approved_amount"],
            row["from_treatment_date"], row["to_treatment_date"],
            row["created_date"], row["limit_id"], row["quarter_id"],
            "YES" if TAG in (row["remark"] or "") else "NO", row["remark"],
        ])
        ws.cell(ws.max_row, 6).number_format = MONEY
        ws.cell(ws.max_row, 7).number_format = MONEY
        if TAG in (row["remark"] or ""):
            for cell in ws[ws.max_row]:
                cell.fill = GOOD_FILL
    style_table(ws, "OutdoorClaimDetails")
    fit_columns(ws, 60)


def add_migration_audit(wb, claims):
    ws = wb.create_sheet("Migration Audit")
    ws.append([
        "Batch Marker", "Action", "EPF", "Claim ID", "Request ID", "Status",
        "Approved Amount", "Final Category", "Quarter ID", "Source / Remark", "Result",
    ])
    tagged = [row for row in claims if TAG in (row["remark"] or "")]
    for row in tagged:
        action, result = AUDIT_ACTIONS.get(row["claim_id"], ("RECONCILED", "Reconciled in migration batch."))
        ws.append([
            TAG, action, str(row["epf_no"]), row["claim_id"], row["request_id"],
            row["request_status"], row["approved_amount"], row["treatment_category"],
            row["quarter_id"], row["remark"], result,
        ])
        ws.cell(ws.max_row, 7).number_format = MONEY
    style_table(ws, "OutdoorMigrationAudit")
    fit_columns(ws, 70)


def add_source_reconciliation(wb, cohort, claims):
    ws = wb.create_sheet("Source Reconciliation")
    ws.append([
        "EPF", "Employee Name", "Expected Approved Total", "Live Approved OTHER",
        "Live Approved DENTAL", "Live Approved Total", "Live Balance", "Match",
    ])
    names = {str(r["epf_no"]): r["employee_name"] for r in cohort}
    by_epf = defaultdict(list)
    for row in claims:
        if row["request_status"] == "APPROVED":
            by_epf[str(row["epf_no"])].append(row)
    for epf, expected in EXPECTED_SOURCE.items():
        rows = by_epf[epf]
        other = sum((r["approved_amount"] or Decimal(0)) for r in rows if r["treatment_category"] == "OTHER")
        dental = sum((r["approved_amount"] or Decimal(0)) for r in rows if r["treatment_category"] == "DENTAL")
        total = other + dental
        match = "YES" if total == expected else "NO"
        ws.append([epf, names.get(epf, ""), expected, other, dental, total, GLOBAL_LIMIT-total, match])
        for col in range(3, 8):
            ws.cell(ws.max_row, col).number_format = MONEY
        ws.cell(ws.max_row, 8).fill = GOOD_FILL if match == "YES" else WARN_FILL
    style_table(ws, "OutdoorSourceReconciliation")
    fit_columns(ws)


def add_limit_configuration(wb, global_row, quarters):
    ws = wb.create_sheet("Limit Configuration")
    ws.append([
        "Limit ID", "Policy Period ID", "Staff Category", "Insurance Policy",
        "Treatment", "Treatment Category", "Configured Limit", "From Date",
        "To Date", "Description",
    ])
    ws.append([
        global_row["limit_id"], global_row["period_id"], "NS",
        global_row["insurance_policy"], global_row["treatment"], "GLOBAL",
        global_row["global_limit"], datetime(2025, 1, 21), datetime(2026, 1, 20),
        "Shared maximum across approved OTHER and DENTAL claims.",
    ])
    for row in quarters:
        description = "Annual dental sublimit" if row["treatment_category_code"] == "DENTAL" else "Date-based OTHER cap"
        ws.append([
            LIMIT_ID, global_row["period_id"], "NS", global_row["insurance_policy"],
            global_row["treatment"], row["treatment_category_code"],
            row["quarter_limit"], row["from_date"], row["to_date"], description,
        ])
    for row_no in range(2, ws.max_row + 1):
        ws.cell(row_no, 7).number_format = MONEY
    style_table(ws, "OutdoorLimitConfiguration")
    fit_columns(ws)


def validate_report(path: Path, expected_claim_count: int):
    wb = load_workbook(path, read_only=True, data_only=True)
    required = {
        "Summary", "Employee Category Limits", "Claim Details",
        "Migration Audit", "Source Reconciliation", "Limit Configuration",
    }
    if set(wb.sheetnames) != required:
        raise RuntimeError(f"Unexpected report sheets: {wb.sheetnames}")
    if wb["Migration Audit"].max_row != 10:
        raise RuntimeError("Migration Audit must contain exactly nine tagged claims")
    if wb["Claim Details"].max_row != expected_claim_count + 1:
        raise RuntimeError("Claim Details row count does not match the live query")
    matches = [row[7] for row in wb["Source Reconciliation"].iter_rows(min_row=2, values_only=True)]
    if not matches or set(matches) != {"YES"}:
        raise RuntimeError(f"Source reconciliation mismatch: {matches}")


def main():
    with closing(_connect()) as connection, closing(connection.cursor(dictionary=True)) as cur:
        db = query_all(cur, "SELECT DATABASE() db")[0]["db"]
        if db != "sgcs_care":
            raise RuntimeError(f"Expected live sgcs_care database, found {db!r}")
        cohort = active_cohort(cur)
        claims = claim_rows(cur, [row["user_id"] for row in cohort])
        global_row, quarters = limit_rows(cur)

    marker_count = sum(TAG in (row["remark"] or "") for row in claims)
    if len(cohort) != 26:
        raise RuntimeError(f"Expected 26 employees including EPF 206 override, found {len(cohort)}")
    if marker_count != 9:
        raise RuntimeError(f"Expected nine tagged migration claims, found {marker_count}")

    wb = Workbook()
    generated_at = datetime.now(ZoneInfo("Asia/Colombo")).strftime("%Y-%m-%d %H:%M:%S Asia/Colombo")
    add_summary(wb, cohort, claims, marker_count, generated_at)
    add_employee_limits(wb, cohort, claims)
    add_claim_details(wb, claims)
    add_migration_audit(wb, claims)
    add_source_reconciliation(wb, cohort, claims)
    add_limit_configuration(wb, global_row, quarters)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    validate_report(OUTPUT, len(claims))
    print(f"Created {OUTPUT.resolve()}")
    print(f"Employees={len(cohort)} Claims={len(claims)} Tagged={marker_count} SourceMatches={len(EXPECTED_SOURCE)}")


if __name__ == "__main__":
    main()
